"""
build_lab_results.py — maintainer tool, participants never need to run this.

Writes ../files/Example_Lab_Results.xlsx: a MOCK laboratory results sheet laid out the way
many soil labs return results (one row per sample: identity, date, masses, N, total C,
inorganic C, organic C, LOI, bulk density, comments), filled with the Cowichan worked-example
values so it lines up with the Example workbook. "Example Lab" is not a real laboratory.
A second tab, "Reading this sheet", maps each column to the digital data sheet.

Usage (from DataAnalysisWorkflow/): python3 data-raw/build_lab_results.py
Recalculate afterwards (LibreOffice headless) so the bulk-density formulas carry values;
data-raw/rebuild_workbooks.sh does this.
"""
import csv
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

OUT = "../files/Example_Lab_Results.xlsx"
DIA = 7.6
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill("solid", fgColor="DCE6F1")
CALC = PatternFill("solid", fgColor="F2F2F2")
NOTE = PatternFill("solid", fgColor="FFF2CC")
# Provenance key: every value on the Results tab is one of these.
PUBLISHED = PatternFill("solid", fgColor="D9EAD3")       # published measurement
RECON = PatternFill("solid", fgColor="FCE5CD")           # reconstructed teaching value
ILLUS = PatternFill("solid", fgColor="E6E0F0")           # illustrative detail
KEY = [(PUBLISHED, "Published measurement (Douglas et al. 2022)"),
       (RECON, "Reconstructed teaching value — dry mass back-calculated from published bulk density"),
       (ILLUS, "Illustrative detail — lab ID, date, comments; not real"),
       (CALC, "Calculated on this sheet")]


def write_key(ws, row, cols=(1, 4, 9, 13)):
    """One row of coloured swatches, so any crop of the table header carries the key."""
    for (fill, text), c in zip(KEY, cols):
        cell = ws.cell(row, c, text)
        cell.fill = fill; cell.border = BOX
        cell.font = Font(size=9, bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 30


def cowichan_rows():
    cores = {r["SampID"]: r for r in csv.DictReader(open("data/reference/janousek2025_zostera_cores.csv"))
             if r["Estuary"] == "COW"}
    depth = [r for r in csv.DictReader(open("data/reference/janousek2025_zostera_depthseries.csv"))
             if r["SampID"] in cores and r["BD_type"] == "M"]
    out = []
    for sid in sorted(cores):
        rows = sorted((d for d in depth if d["SampID"] == sid), key=lambda d: float(d["depth_top_cm"]))
        for k, d in enumerate(rows, start=1):
            top, bot = float(d["depth_top_cm"]), float(d["depth_bottom_cm"])
            area = 3.141592653589793 * (DIA / 2) ** 2
            out.append(dict(core=f"COW-{d['StudySampID']}", sample=k, top=top, bot=bot,
                            dry=round(float(d["BD"]) * area * (bot - top), 2),
                            oc=float(d["PercC"]) if d["C_type"] == "M" and d["PercC"] not in ("", "NA") else None,
                            loi=float(d["PercOM"]) if d["PercOM"] not in ("", "NA") else None))
    return out


def build():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    ws["A1"] = "Example Lab — sediment carbon results"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("MOCK RESULTS FOR TEACHING. 'Example Lab' is not a real laboratory. Values are the Cowichan "
                "worked example (Douglas et al. 2022, via Janousek et al. 2025), laid out the way a lab "
                "might return them.")
    ws["A2"].font = Font(italic=True, color="C00000")
    meta = [("Client", "Blue Carbon Eelgrass Workshop (example)"), ("Samples received", "45 slices from 3 cores"),
            ("Drying (illustrative)", "65 °C to constant mass"), ("Loss on ignition", "550 °C, 5 h"),
            ("Organic carbon", "Elemental analyser after acid fumigation (carbonate removed). Run on a subset."),
            ("Bulk density", f"Dry mass ÷ (π × (corer internal diameter ÷ 2)² × slice thickness); "
                             f"corer internal diameter {DIA} cm as supplied by the client")]
    for i, (a, b) in enumerate(meta, start=4):
        ws.cell(i, 1, a).font = Font(bold=True)
        ws.cell(i, 3, b)
    hdr = ["Lab ID", "Client sample ID", "Date received", "Top (cm)", "Bottom (cm)",
           "Slice thickness (cm)", "Corer internal diameter (cm)", "Dry mass for BD (g)",
           "Dry bulk density (g/cm³)", "N (%)", "C total (%)", "C inorganic (%)", "C organic (%)",
           "LOI 550 °C (%)", "Comments"]
    h = 11
    for c, t in enumerate(hdr, start=1):
        cell = ws.cell(h, c, t)
        cell.font = Font(bold=True); cell.fill = HEAD; cell.border = BOX
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[h].height = 60
    write_key(ws, h - 1)
    ws.merge_cells(start_row=h - 1, start_column=1, end_row=h - 1, end_column=3)
    ws.merge_cells(start_row=h - 1, start_column=4, end_row=h - 1, end_column=8)
    ws.merge_cells(start_row=h - 1, start_column=9, end_row=h - 1, end_column=12)
    ws.merge_cells(start_row=h - 1, start_column=13, end_row=h - 1, end_column=15)
    # provenance of each column: P published, R reconstructed, I illustrative, C calculated
    prov = {1: ILLUS, 2: PUBLISHED, 3: ILLUS, 4: PUBLISHED, 5: PUBLISHED, 6: CALC, 7: PUBLISHED, 8: RECON,
            9: CALC, 13: PUBLISHED, 14: PUBLISHED, 15: ILLUS}
    for i, r in enumerate(cowichan_rows()):
        row = h + 1 + i
        vals = [f"EL-{1001 + i}", f"{r['core']}-{r['sample']:02d}", "2026-08-04", r["top"], r["bot"],
                f"=E{row}-D{row}", DIA, r["dry"],
                f"=IF(OR(F{row}=\"\",G{row}=\"\",H{row}=\"\"),\"\",H{row}/(PI()*(G{row}/2)^2*F{row}))",
                None, None, None, r["oc"], r["loi"],
                "OC run" if r["oc"] is not None else "LOI only (OC not requested for this slice)"]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row, c, v)
            cell.border = BOX
            if c in prov and v is not None:
                cell.fill = prov[c]
        ws.cell(row, 9).number_format = "0.000"
        ws.cell(row, 8).number_format = "0.00"
    last = h + len(cowichan_rows())
    ws.cell(last + 2, 1, "N, C total and C inorganic were not run: organic carbon was measured directly after "
                         "acid fumigation. A blank means 'not measured' — never zero.").font = Font(italic=True)
    for c, w in enumerate([9, 15, 12, 8, 9, 10, 12, 11, 12, 7, 9, 10, 10, 10, 38], start=1):
        ws.column_dimensions[L(c)].width = w
    ws.freeze_panes = f"C{h + 1}"

    # ---------------------------------------------------------------- Reading this sheet
    rd = wb.create_sheet("Reading this sheet")
    rd["A1"] = "Reading a lab results sheet — and moving it into the digital data sheet"
    rd["A1"].font = Font(bold=True, size=13)
    rd["A2"] = "Each row of the Results tab is one slice. Match it to Sheet 3 of the digital data sheet by Core ID and Sample ID."
    write_key(rd, 3, cols=(1, 2, 3, 4))
    rows = [
        ("On the results sheet", "Goes to (digital data sheet, Sheet 3)", "Check before you copy"),
        ("Client sample ID (e.g. COW-S5-01)", "Column A Core ID (COW-S5) and column B Sample ID (1)",
         "The ID must be exactly what you wrote on the bags and the field data sheet."),
        ("Dry mass for BD (g)", "Column M — Dry weight (g), whole slice",
         "Grams, not kilograms. It must be the dry mass of the WHOLE slice, not of a subsample."),
        ("C organic (%)", "Column N — Carbon value, and type OC in column O",
         "Only if inorganic carbon was removed or subtracted. A percent of dry mass (0.9 means 0.9%)."),
        ("LOI 550 °C (%)", "Column N — Carbon value, and type LOI in column O (only where there is no C organic)",
         "Needs your LOI→OC equation on the Instructions tab. Note the temperature and duration."),
        ("C total (%)", "Not used as organic carbon", "Total carbon includes shell carbonate. Ask the lab for inorganic carbon, then OC = TC − IC."),
        ("C inorganic (%)", "Not entered", "Used only to work out OC = TC − IC."),
        ("N (%)", "Not entered", "Useful for C:N (a source indicator), not needed for the stock."),
        ("Dry bulk density (g/cm³)", "Not entered — the data sheet calculates its own (column Q)",
         "Use the lab's value only as a cross-check. It must divide by slice thickness: dry mass ÷ (π r² × thickness)."),
        ("Drying temperature", "Note it in the workbook (Sheet 3 notes or Sheet 2 core notes)",
         "60–65 °C is usual for carbon. Material dried at 105 °C can lose some organic matter."),
    ]
    for i, (a, b, c) in enumerate(rows, start=5):
        for j, v in enumerate((a, b, c), start=1):
            cell = rd.cell(i, j, v)
            cell.border = BOX
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if i == 5:
                cell.font = Font(bold=True); cell.fill = HEAD
    n = 5 + len(rows) + 1
    rd.cell(n, 1, "Three things to look for on any results sheet").font = Font(bold=True)
    tips = ["1. Bulk density divided by slice thickness. If the lab divides dry mass by the corer's cross-section "
            "only (π r²), a 2 cm slice comes out twice as dense as it is. Check the formula whenever values look "
            "unusual for your sediment — dense sands can genuinely exceed 1.8 g/cm³ (several Cowichan slices reach 2.0).",
            "2. Units. Masses in g or kg, carbon as % or as a fraction (0–1), and whether the units changed between reports.",
            "3. What 'C' means. Total carbon is not organic carbon. If the sheet does not say inorganic carbon was removed, ask."]
    for k, t in enumerate(tips, start=n + 1):
        cell = rd.cell(k, 1, t)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.fill = NOTE
        rd.merge_cells(start_row=k, start_column=1, end_row=k, end_column=3)
        rd.row_dimensions[k].height = 32
    for c, w in enumerate([34, 52, 70, 30], start=1):
        rd.column_dimensions[L(c)].width = w
    wb.save(OUT)
    print("written", OUT)


if __name__ == "__main__":
    build()
