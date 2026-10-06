"""
build_workbooks.py — maintainer tool, participants never need to run this.

Rebuilds the workshop's digital data sheet from the original 2026 template, keeping its
layout, colours and tab names, and applying the agreed corrections:

  * Sample Data gains a "Carbon measured as" column (OC / TC / LOI) so organic carbon,
    total carbon and loss-on-ignition are never treated as the same thing; LOI is converted
    with a stated equation (Instructions tab); TC is refused until inorganic carbon is removed.
  * Each slice's stock is split across the standard depth increments (0-15, 15-30, 30-50,
    50-100 cm, in-situ depths) in proportion to overlap — mass-conserving.
  * A Slice check column flags duplicates, overlaps, gaps, missing lab values and TC.
  * Core Summary only totals a core when every slice is checked OK — a blank is never a zero —
    and reports increments only where the core reaches them.
  * Plot & Core Log gains Stratum, an explicit "compaction not measured" choice, and Core notes.

Usage (from DataAnalysisWorkflow/):
  python3 data-raw/build_workbooks.py <original_template.xlsx>
Writes ../files/Eelgrass_Carbon_DigitalData_BlankSheet.xlsx and
       ../files/Eelgrass_Carbon_DigitalData_Example.xlsx (Cowichan Estuary example).
Recalculate afterwards (LibreOffice headless) so cached values exist for readers.
"""
import csv, math, sys
from copy import copy
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as L

TEMPLATE = sys.argv[1]
OUT_DIR = "../files"
N_CORE_ROWS = (6, 35)      # Plot & Core Log rows
N_SLICE_ROWS = (6, 205)    # Sample Data rows
N_SUM_ROWS = (5, 34)       # Core Summary rows
INCREMENTS = [(0, 15), (15, 30), (30, 50), (50, 100)]
S2, S3, S4 = "'2. Plot & Core Log'", "'3. Sample Data'", "'4. Core Summary'"


def style_of(cell):
    return dict(font=copy(cell.font), fill=copy(cell.fill), border=copy(cell.border),
                alignment=copy(cell.alignment), number_format=cell.number_format)


def apply(cell, st, number_format=None):
    cell.font, cell.fill, cell.border = copy(st["font"]), copy(st["fill"]), copy(st["border"])
    cell.alignment = copy(st["alignment"])
    cell.number_format = number_format or st["number_format"]


def unmerge_all(ws):
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))


def clear(ws, rows, cols):
    for r in range(rows[0], rows[1] + 1):
        for c in range(1, cols + 1):
            ws.cell(r, c).value = None


# ----------------------------------------------------------------------------- build
def build(example=None):
    wb = openpyxl.load_workbook(TEMPLATE)
    ins, log, smp, summ = (wb[n] for n in
                           ["1. Instructions", "2. Plot & Core Log", "3. Sample Data", "4. Core Summary"])

    # ---- captured styles (from the original template, so the look is unchanged)
    st = {
        "title": style_of(smp["A1"]), "sub": style_of(smp["A2"]),
        "band_field": style_of(smp["A4"]), "band_pre": style_of(smp["F4"]),
        "band_lab": style_of(smp["L4"]), "band_post": style_of(smp["O4"]),
        "hdr_field": style_of(smp["A5"]), "hdr_pre": style_of(smp["F5"]),
        "hdr_lab": style_of(smp["L5"]), "hdr_post": style_of(smp["O5"]),
        "input": style_of(smp["A6"]), "calc": style_of(smp["F6"]), "lookup": style_of(smp["G6"]),
        "sum_hdr": style_of(summ["A4"]),
        "i_head": style_of(ins["A5"]), "i_label": style_of(ins["A6"]), "i_text": style_of(ins["B6"]),
        "i_title": style_of(ins["A1"]), "i_sub": style_of(ins["A2"]),
    }
    lab_input = style_of(smp["L6"])

    # =========================================================== 2. Plot & Core Log
    unmerge_all(log)
    for c in range(16, 19):
        apply(log.cell(4, c), st["band_field"])
    for c in range(19, 21):
        apply(log.cell(4, c), st["band_pre"])
    log["P4"] = "Design & notes"
    log["S4"] = "Corer used"
    heads = {16: "Stratum (code)\ne.g. SG", 17: "Compaction not measured?\nenter: assume none",
             18: "Core notes"}
    for c, t in heads.items():
        log.cell(5, c).value = t
        apply(log.cell(5, c), style_of(log["A5"]))
    for c, t in {19: "Corer diameter\nused (cm)", 20: "Diameter\nfrom"}.items():
        log.cell(5, c).value = t
        apply(log.cell(5, c), style_of(log["M5"]))
    log["J5"] = "Corer internal\ndiameter (cm)\nonly if different\nfrom Instructions"
    log["K5"] = "Outside depth (cm)\n= depth of corer inserted"
    log["L5"] = "Inside depth (cm)\n= length of core extracted"
    c0, c1 = N_CORE_ROWS
    for r in range(c0, c1 + 1):
        for c in (16, 17, 18):
            apply(log.cell(r, c), st["input"])
        for c in (19, 20):
            apply(log.cell(r, c), style_of(log[f"M{r}"]), "0.00" if c == 19 else "General")
        log[f"M{r}"] = (f'=IF(AND($K{r}<>"",$L{r}<>""),$K{r}/$L{r},'
                        f'IF(LOWER(TRIM($Q{r}))="assume none",1,""))')
        log[f"N{r}"] = f'=IF(OR($K{r}="",$L{r}=""),"",$K{r}-$L{r})'
        log[f"S{r}"] = f'=IF($B{r}="","",IF($J{r}<>"",$J{r},IF(CORER_DIAMETER_CM<>"",CORER_DIAMETER_CM,"")))'
        log[f"T{r}"] = (f'=IF($B{r}="","",IF($J{r}<>"","this core",'
                        f'IF(CORER_DIAMETER_CM<>"","Instructions","missing")))')
        log[f"O{r}"] = (
            f'=IF($B{r}="","",IF(COUNTIF($B${c0}:$B${c1},$B{r})>1,"CHECK: duplicate Core ID",'
            f'IF($S{r}="","CHECK: corer diameter missing",'
            f'IF(AND($K{r}="",$L{r}=""),IF($M{r}=1,"ASSUMED: no compaction (not measured)",'
            f'"CHECK: compaction not recorded"),IF(OR($K{r}="",$L{r}=""),"CHECK: one depth missing",'
            f'IF($L{r}>$K{r},"CHECK: extracted > inserted",IF($M{r}>1.5,"CHECK: >50% compaction","OK")))))))')
    log.column_dimensions["P"].width = 12
    log.column_dimensions["Q"].width = 20
    log.column_dimensions["R"].width = 34
    log.column_dimensions["S"].width = 11
    log.column_dimensions["T"].width = 11
    log.merge_cells("A2:T2"); log.merge_cells("A4:I4"); log.merge_cells("J4:L4")
    log.merge_cells("M4:O4"); log.merge_cells("P4:R4"); log.merge_cells("S4:T4"); log.merge_cells("A38:T39")
    log["A2"] = ("Enter one row for every core you collect. Core ID must be unique and must match "
                 "exactly what you type on Sheet 3 — the sample rows look up their compaction factor "
                 "and corer diameter from here.")
    log["A38"] = (
        "Outside depth = how far the corer was driven into the sediment (the datasheet's 'Depth of corer "
        "inserted'). Inside depth = the length of core actually recovered in the tube ('Length of core "
        "extracted'). Compaction factor = outside / inside: the number a measured depth is multiplied by "
        "to give its true in-situ depth. If compaction was not measured, leave both blank and type "
        "'assume none' in column Q only if you have a reason to believe there was none — the QC check "
        "will then say ASSUMED, never OK. Stratum: the code from your sampling design (Part 2), e.g. SG; "
        "leave blank if you did not stratify. Corer diameter: enter it once on the Instructions tab (YOUR "
        "CORER); fill column J only for a core taken with a different tube. Column S shows the diameter used.")
    log.row_dimensions[38].height = 30; log.row_dimensions[39].height = 30
    log["A38"].alignment = openpyxl.styles.Alignment(
        wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"assume none"', allow_blank=True)
    log.add_data_validation(dv); dv.add(f"Q{N_CORE_ROWS[0]}:Q{N_CORE_ROWS[1]}")

    # =========================================================== 3. Sample Data
    unmerge_all(smp)
    old_widths = {k: v.width for k, v in smp.column_dimensions.items()}
    clear(smp, (4, 207), 30)
    cols = [  # (header, band, kind, number format)
        ("Core ID", "field", "input", None), ("Sample ID", "field", "input", None),
        ("Top depth\n(cm)", "field", "input", "0.0"), ("Bottom depth\n(cm)", "field", "input", "0.0"),
        ("Notes", "field", "input", None),
        ("Depth interval\n(cm)", "pre", "calc", "0.0"), ("Compaction\nfactor", "pre", "lookup", "0.000"),
        ("Corer diameter\n(cm)", "pre", "lookup", "0.00"), ("In-situ top\ndepth (cm)", "pre", "calc", "0.0"),
        ("In-situ bottom\ndepth (cm)", "pre", "calc", "0.0"), ("Sample volume\n(cm3)", "pre", "calc", "0.0"),
        ("Wet weight\n(g)", "lab", "labin", "0.00"), ("Dry weight\n(g)\nwhole slice", "lab", "labin", "0.00"),
        ("Carbon value\n(% of dry mass)", "lab", "labin", "0.00"),
        ("Carbon measured as\nOC / TC / LOI", "lab", "labin", None),
        ("Water content\n(%)", "post", "calc", "0.0"),
        ("Dry bulk density\n(g/cm3)", "post", "calc", "0.000"),
        ("Organic carbon\n(%)", "post", "calc", "0.000"), ("Organic carbon\n(g/kg)", "post", "calc", "0.00"),
        ("Carbon density\n(g C/cm3)", "post", "calc", "0.00000"),
        ("Carbon stock\n(g C/cm2)", "post", "calc", "0.0000"), ("Carbon stock\n(kg C/m2)", "post", "calc", "0.000"),
    ] + [(f"Stock in\n{a}–{b} cm\n(kg C/m2)", "post", "calc", "0.000") for a, b in INCREMENTS] + [
        ("Slice check", "check", "calc", None)]
    band_ranges = {"field": (1, 5, "1. Data from the field"), "pre": (6, 11, "2. Calculated before the lab"),
                   "lab": (12, 15, "3. Measured by the lab"), "post": (16, 26, "4. Calculated after the lab"),
                   "check": (27, 27, "5. Check")}
    band_style = {"field": "band_field", "pre": "band_pre", "lab": "band_lab", "post": "band_post",
                  "check": "band_field"}
    hdr_style = {"field": "hdr_field", "pre": "hdr_pre", "lab": "hdr_lab", "post": "hdr_post", "check": "hdr_field"}
    for key, (c0, c1, label) in band_ranges.items():
        for c in range(c0, c1 + 1):
            apply(smp.cell(4, c), st[band_style[key]])
        smp.cell(4, c0).value = label
        if c1 > c0:
            smp.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c1)
    for i, (h, band, kind, nf) in enumerate(cols, start=1):
        smp.cell(5, i).value = h
        apply(smp.cell(5, i), st[hdr_style[band]])
    widths = [13, 11, 10, 11, 30, 11, 11, 12, 11, 12, 11, 11, 12, 13, 15, 12, 14, 13, 13, 13, 12, 12,
              12, 12, 12, 12, 30]
    for i, w in enumerate(widths, start=1):
        smp.column_dimensions[L(i)].width = w
    lk = lambda col, r: (f'=IF($A{r}="","",IFERROR(INDEX({S2}!${col}${N_CORE_ROWS[0]}:${col}${N_CORE_ROWS[1]},'
                         f'MATCH($A{r},{S2}!$B${N_CORE_ROWS[0]}:$B${N_CORE_ROWS[1]},0)),""))')
    a0, a1 = N_SLICE_ROWS
    A, C, D = (f"${x}${a0}:${x}${a1}" for x in "ACD")
    # inside depth (length extracted) of this slice's core; "" when not recorded. COUNTIFS/SUMIFS,
    # not INDEX, because INDEX of a blank input cell returns 0 rather than "".
    rB, rL = (f"{S2}!${x}${N_CORE_ROWS[0]}:${x}${N_CORE_ROWS[1]}" for x in "BL")
    core_len = lambda r: f'IF(COUNTIFS({rB},$A{r},{rL},">0")=0,"",SUMIFS({rL},{rB},$A{r}))'
    for r in range(a0, a1 + 1):
        f = {
            6: f'=IF(OR($C{r}="",$D{r}=""),"",$D{r}-$C{r})',
            7: lk("M", r), 8: lk("S", r),
            9: f'=IF(OR($C{r}="",$G{r}=""),"",$C{r}*$G{r})',
            10: f'=IF(OR($D{r}="",$G{r}=""),"",$D{r}*$G{r})',
            11: f'=IF(OR($F{r}="",$H{r}=""),"",PI()*(($H{r}/2)^2)*$F{r})',
            16: f'=IF(OR($L{r}="",$M{r}="",$L{r}=0),"",(($L{r}-$M{r})/$L{r})*100)',
            17: f'=IF(OR($M{r}="",$K{r}="",$K{r}=0),"",$M{r}/$K{r})',
            18: (f'=IF(OR($N{r}="",$O{r}=""),"",IF(UPPER($O{r})="OC",$N{r},IF(UPPER($O{r})="LOI",'
                 f'IF(OR(LOI_INTERCEPT="",LOI_SLOPE=""),"",IF(LOI_INTERCEPT+LOI_SLOPE*$N{r}<0,"",'
                 f'LOI_INTERCEPT+LOI_SLOPE*$N{r})),"")))'),
            19: f'=IF($R{r}="","",$R{r}*10)',
            20: f'=IF(OR($Q{r}="",$R{r}=""),"",$Q{r}*$R{r}/100)',
            21: f'=IF(OR($T{r}="",$F{r}=""),"",$T{r}*$F{r})',
            22: f'=IF($U{r}="","",$U{r}*10)',
        }
        for k, (a, b) in enumerate(INCREMENTS):
            f[23 + k] = (f'=IF(OR($V{r}="",$I{r}="",$J{r}="",$J{r}<=$I{r}),"",'
                         f'$V{r}*MAX(0,MIN($J{r},{b})-MAX($I{r},{a}))/($J{r}-$I{r}))')
        f[27] = (
            f'=IF($A{r}="","",IF(OR($C{r}="",$D{r}="",$D{r}<=$C{r}),"CHECK: depths missing, or bottom not below top",'
            f'IF(COUNTIFS({A},$A{r},{C},$C{r})>1,"DUPLICATE slice",'
            f'IF(COUNTIFS({A},$A{r},{C},"<"&$D{r},{D},">"&$C{r})>1,"OVERLAP with another slice",'
            f'IF(AND($C{r}>0,COUNTIFS({A},$A{r},{D},$C{r})=0),"GAP above this slice",'
            f'IF(OR($G{r}="",$H{r}=""),"CHECK: Core ID not in Core Log, or its compaction / diameter missing",'
            f'IF(AND({core_len(r)}<>"",$D{r}>{core_len(r)}),"CHECK: slice deeper than the core length on Sheet 2",'
            f'IF(OR($M{r}="",$N{r}="",$O{r}=""),"AWAITING LAB (dry weight, carbon value or type)",'
            f'IF(AND(UPPER($O{r})<>"OC",UPPER($O{r})<>"LOI"),IF(UPPER($O{r})="TC",'
            f'"TC is not organic carbon — ask the lab for OC or IC","CHECK: type must be OC, TC or LOI"),'
            f'IF($R{r}="",IF(AND(LOI_INTERCEPT<>"",LOI_SLOPE<>""),'
            f'"CHECK: LOI below the range of the conversion equation (OC would be < 0)",'
            f'"CHECK: LOI conversion not set (Instructions tab)"),'
            f'IF(OR($Q{r}<=0,$Q{r}>2.65,$R{r}>50),"CHECK: bulk density outside 0–2.65 g/cm3 or OC > 50%",'
            f'"OK"))))))))))')
        for i, (h, band, kind, nf) in enumerate(cols, start=1):
            cell = smp.cell(r, i)
            apply(cell, {"input": st["input"], "labin": lab_input, "calc": st["calc"],
                         "lookup": st["lookup"]}[kind], nf)
            if i in f:
                cell.value = f[i]
    smp.merge_cells("A2:AA2")
    smp["A1"] = "Sample Data — one row per slice"
    smp["A2"] = ("Type only in the yellow columns. Core ID must match Sheet 2 exactly. Enter the carbon value "
                 "as a percent of dry mass (0.9 means 0.9%), and say what it is in column O: OC, TC or LOI.")
    smp["A207"] = ("Depths in columns C–D are measured down the recovered core (in the tube). The increment "
                   "columns W–Z use in-situ depths (I–J), splitting a slice that crosses a boundary in "
                   "proportion to its overlap, so no carbon is gained or lost.")
    dv = DataValidation(type="list", formula1='"OC,TC,LOI"', allow_blank=True)
    smp.add_data_validation(dv); dv.add(f"O{a0}:O{a1}")
    smp.freeze_panes = "F6"

    # =========================================================== 4. Core Summary
    unmerge_all(summ)
    clear(summ, (4, 40), 20)
    sh = ["Core ID", "Plot ID", "Stratum", "Slices entered", "Slices checked OK",
          "Measured to —\nin tube (cm)", "Measured to —\nin situ (cm)", "Status",
          "Dry bulk density\n(g/cm3)\nthickness-weighted", "Organic carbon\n(%)\nmass-weighted",
          "Measured core stock\n(kg C/m2)", "Measured core stock\n(Mg C/ha)"] + \
         [f"{a}–{b} cm\n(kg C/m2)" for a, b in INCREMENTS]
    for i, h in enumerate(sh, start=1):
        summ.cell(4, i).value = h
        apply(summ.cell(4, i), st["sum_hdr"])
    for i, w in enumerate([14, 10, 9, 10, 11, 13, 13, 34, 15, 14, 15, 15, 14, 14, 14, 14], start=1):
        summ.column_dimensions[L(i)].width = w
    rngS = lambda col: f"{S3}!${col}${a0}:${col}${a1}"
    dia = lambda r: (f"INDEX({S2}!$S${N_CORE_ROWS[0]}:$S${N_CORE_ROWS[1]},"
                     f"MATCH($A{r},{S2}!$B${N_CORE_ROWS[0]}:$B${N_CORE_ROWS[1]},0))")
    for k, r in enumerate(range(N_SUM_ROWS[0], N_SUM_ROWS[1] + 1)):
        lr = N_CORE_ROWS[0] + k
        f = {
            1: f"=IF({S2}!$B${lr}=\"\",\"\",{S2}!$B${lr})",
            2: f"=IF($A{r}=\"\",\"\",IF({S2}!$A${lr}=\"\",\"\",{S2}!$A${lr}))",
            3: f"=IF($A{r}=\"\",\"\",IF({S2}!$P${lr}=\"\",\"\",{S2}!$P${lr}))",
            4: f'=IF($A{r}="","",COUNTIFS({rngS("A")},$A{r}))',
            5: f'=IF($A{r}="","",COUNTIFS({rngS("A")},$A{r},{rngS("AA")},"OK"))',
            6: f'=IF(OR($A{r}="",$D{r}=0),"",_xlfn.MAXIFS({rngS("D")},{rngS("A")},$A{r}))',
            7: f'=IF(OR($A{r}="",$D{r}=0),"",IFERROR(_xlfn.MAXIFS({rngS("J")},{rngS("A")},$A{r}),""))',
            8: (f'=IF($A{r}="","",IF($D{r}=0,"No slices entered",'
                f'IF(AND({S2}!$O${lr}<>"OK",LEFT({S2}!$O${lr},7)<>"ASSUMED"),{S2}!$O${lr}&" (Sheet 2)",'
                f'IF($E{r}<$D{r},"Not complete — see Slice check on Sheet 3","Complete"))))'),
            9: (f'=IF(OR($A{r}="",$H{r}<>"Complete"),"",SUMIFS({rngS("M")},{rngS("A")},$A{r})/'
                f'(PI()*({dia(r)}/2)^2*SUMIFS({rngS("F")},{rngS("A")},$A{r})))'),
            10: (f'=IF(OR($A{r}="",$H{r}<>"Complete"),"",100*SUMIFS({rngS("U")},{rngS("A")},$A{r})/'
                 f'(SUMIFS({rngS("M")},{rngS("A")},$A{r})/(PI()*({dia(r)}/2)^2)))'),
            11: f'=IF(OR($A{r}="",$H{r}<>"Complete"),"",SUMIFS({rngS("V")},{rngS("A")},$A{r}))',
            12: f'=IF($K{r}="","",$K{r}*10)',
        }
        for j, (a, b) in enumerate(INCREMENTS):
            col = L(23 + j)
            f[13 + j] = (f'=IF(OR($A{r}="",$H{r}<>"Complete"),"",IF($G{r}>={b}-0.0001,'
                         f'SUMIFS({rngS(col)},{rngS("A")},$A{r}),IF($G{r}>{a},'
                         f'"partial — to "&TEXT($G{r},"0.0")&" cm","not reached")))')
        nfs = {6: "0.0", 7: "0.0", 9: "0.000", 10: "0.000", 11: "0.000", 12: "0.0",
               13: "0.000", 14: "0.000", 15: "0.000", 16: "0.000"}
        for i in range(1, 17):
            apply(summ.cell(r, i), st["lookup"] if i <= 3 else st["calc"], nfs.get(i, "General"))
            summ.cell(r, i).value = f[i]
    mr = N_SUM_ROWS[1] + 2
    summ.cell(mr, 1).value = "Mean of complete cores"
    summ.cell(mr + 1, 1).value = "Number of cores in that mean"
    for i in range(1, 17):
        apply(summ.cell(mr, i), st["calc"], "0.000"); apply(summ.cell(mr + 1, i), st["calc"], "0")
    summ.cell(mr, 1).font = copy(summ.cell(4, 1).font); summ.cell(mr, 1).fill = copy(summ.cell(4, 1).fill)
    summ.cell(mr + 1, 1).font = copy(summ.cell(4, 1).font); summ.cell(mr + 1, 1).fill = copy(summ.cell(4, 1).fill)
    for i in range(13, 17):
        c = L(i)
        summ.cell(mr, i).value = f'=IF(COUNT({c}{N_SUM_ROWS[0]}:{c}{N_SUM_ROWS[1]})=0,"",AVERAGE({c}{N_SUM_ROWS[0]}:{c}{N_SUM_ROWS[1]}))'
        summ.cell(mr + 1, i).value = f'=COUNT({c}{N_SUM_ROWS[0]}:{c}{N_SUM_ROWS[1]})'
    summ.cell(mr + 2, 1).value = (
        "Each increment's mean uses only the cores that reach the bottom of that increment, so the "
        "number of cores can differ between columns. 'Measured core stock' is to whatever depth each core "
        "reached, so it is NOT comparable between cores of different lengths — compare increments instead. "
        "A mean of cores is not an estimate for a meadow: see Part 4, Option B.")
    summ.merge_cells(start_row=mr + 2, start_column=1, end_row=mr + 4, end_column=16)
    summ.cell(mr + 2, 1).alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
    summ.merge_cells("A2:P2")
    summ["A1"] = "Core Summary — automatic, nothing to type"
    summ["A2"] = ("One row per core on Sheet 2. A core is totalled only when its QC check on Sheet 2 is OK (or ASSUMED) "
                  "and every one of its slices passes the Slice check on Sheet 3 — a missing value is never counted as "
                  "zero. Increments use in-situ depths.")
    summ.freeze_panes = "B5"

    # =========================================================== 1. Instructions
    clear(ins, (5, 80), 3)
    for r in range(5, 81):
        for c in (1, 2):
            apply(ins.cell(r, c), st["i_text"])
    rows = []
    H = lambda t: rows.append(("H", t, None))
    R = lambda a, b: rows.append(("R", a, b))
    B = lambda: rows.append(("B", None, None))
    H("HOW TO USE THIS WORKBOOK")
    R("Sheet 2 — Plot & Core Log", "One row per CORE. Enter plot/core notes, the two compaction depths and (if you "
      "stratified) the stratum. The compaction factor calculates itself. Enter a corer diameter there only for a core "
      "taken with a different tube from the one below.")
    R("Sheet 3 — Sample Data", "One row per SLICE. Enter the field columns before you leave site, then add the lab "
      "columns when results come back. Everything else calculates, and the Slice check column tells you what is "
      "missing or inconsistent.")
    R("Sheet 4 — Core Summary", "Per-core totals and standard depth increments. Fully automatic — nothing to type here.")
    B(); H("YOUR CORER")
    R("Corer internal diameter (cm)", None)
    R("About the diameter", "Measured inside the tube with calipers — not the nominal pipe size (see THINGS THAT "
      "CATCH PEOPLE OUT). Used for every core, unless Sheet 2 column J gives a different value for that core. Leave "
      "it blank and the slices are flagged: the volume, and so the bulk density, cannot be calculated.")
    B(); H("COLOUR KEY")
    R("Yellow fill / blue text", "You type here: data from the field data sheet, or from the lab.")
    R("Grey fill / black text", "Calculated. Do not type over these — you will break the column.")
    R("Green text", "Pulled from another sheet in this workbook.")
    B(); H("THE FIVE BANDS ON SHEET 3")
    R("1. Data from the field", "Copy directly from the field data sheet.")
    R("2. Calculated before the lab", "Depth interval, compaction-corrected (in-situ) depths, and sample volume. You "
      "need the volume to hand to the lab, or to compute bulk density yourself.")
    R("3. Measured by the lab", "Wet weight, dry weight of the whole slice, the carbon value, and what that carbon "
      "value is (column O: OC, TC or LOI).")
    R("4. Calculated after the lab", "Bulk density, organic carbon, carbon density, carbon stock per slice, and the "
      "share of each slice's stock that falls in each standard depth increment.")
    R("5. Check", "One plain-language message per slice. A core is only totalled when all its slices say OK.")
    B(); H("CARBON: WHAT THE LAB MEASURED")
    R("OC", "Organic carbon measured directly — elemental analyser on acid-treated sample, or total carbon minus "
      "measured inorganic carbon. Used as is.")
    R("TC", "Total carbon. Includes shell and other carbonate, which is not organic carbon. It is not used: ask the "
      "lab for inorganic carbon (IC) and enter OC = TC − IC instead.")
    R("LOI", "Loss on ignition (% organic matter). Converted to organic carbon with the equation below: "
      "OC% = intercept + slope × LOI%. A value that would convert to below zero is outside the equation's range "
      "and is flagged, not set to zero.")
    R("LOI intercept", None)
    R("LOI slope", None)
    R("Where the equation comes from", "Best: your own calibration — run elemental analysis on a subset of the same "
      "samples and fit OC against LOI. Otherwise use a published seagrass equation and cite it. Leave these blank "
      "and LOI rows stay unconverted (the Slice check says so). Report the equation you used.")
    B(); H("UNITS")
    R("Depths", "centimetres (cm)")
    R("Weights", "grams (g). Some labs report kilograms — convert before typing.")
    R("Volume", "cubic centimetres (cm3)")
    R("Carbon value", "percent of dry mass (0–100). If the lab reports a fraction (0–1), multiply by 100.")
    R("Dry bulk density", "g/cm3 = dry weight of the whole slice ÷ slice volume")
    R("Carbon stock", "g C/cm2 per slice (field guide Eq 4) and kg C/m2 (Eq 6). × 10 gives Mg C/ha.")
    B(); H("STANDARD DEPTH INCREMENTS")
    R("0–15, 15–30, 30–50, 50–100 cm", "Measured on in-situ depths. A slice that crosses a boundary is split in "
      "proportion to how much of it lies on each side. Sheet 4 only reports an increment when the core reaches "
      "its bottom; otherwise it says 'partial' or 'not reached'.")
    B(); H("THE ANALYSIS (Part 4)")
    R("No export needed", "The R workflow reads this workbook directly. Keep the tab names and the header rows as they "
      "are, save as .xlsx (in Google Sheets: File → Download → Microsoft Excel), and point the workflow's settings "
      "file at it.")
    B(); H("THINGS THAT CATCH PEOPLE OUT")
    R("Measure your actual corer diameter", "Nominal pipe size is not internal diameter. A '3 inch' Schedule 40 PVC "
      "pipe has an internal diameter noticeably larger than 7.62 cm. Volume sits in the denominator of bulk density, "
      "so an assumed diameter biases EVERY carbon stock in the dataset. Measure the ID with calipers and enter it "
      "under YOUR CORER above (and on the field data sheet).")
    R("Compaction corrects depths, not stocks", "Carbon stock per slice uses the MEASURED interval and the recovered-"
      "slice volume. The dry mass in the tube already came from a taller in-situ column, so multiplying by the "
      "compaction factor as well would double-count it. The corrected depths (columns I–J) tell you which in-situ "
      "depth each slice represents, which is what the depth increments need.")
    R("Dry weight and volume must be the same material", "Bulk density needs the dry weight of the WHOLE slice. If the "
      "lab only weighs a subsample, use the whole-slice dry weight you recorded before subsampling — never the "
      "subsample mass against the whole-slice volume.")
    R("A blank is not a zero", "Missing slices, gaps between slices and missing lab values are flagged, and the core "
      "is not totalled until they are resolved. Nothing is filled in for you.")
    r = 5
    loi_rows = {}
    for kind, a, b in rows:
        if kind == "H":
            ins.cell(r, 1).value = a; apply(ins.cell(r, 1), st["i_head"])
        elif kind == "R":
            ins.cell(r, 1).value = a; apply(ins.cell(r, 1), st["i_label"])
            if b is not None:
                ins.cell(r, 2).value = b; apply(ins.cell(r, 2), st["i_text"])
                ins.row_dimensions[r].height = max(15.75, 13.5 * math.ceil(len(b) / 105))
            if a in ("LOI intercept", "LOI slope", "Corer internal diameter (cm)"):
                loi_rows[a] = r
                apply(ins.cell(r, 2), st["input"], "0.000" if a.startswith("LOI") else "0.00")
                ins.cell(r, 2).alignment = openpyxl.styles.Alignment(horizontal="left")
        r += 1
    ins["A2"] = ("Companion to WWF-Canada, Measuring Carbon in Coastal Sediments (2026), and Part 4 of the "
                 "Blue Carbon Eelgrass Workshop.")
    for name, key in (("LOI_INTERCEPT", "LOI intercept"), ("LOI_SLOPE", "LOI slope"),
                      ("CORER_DIAMETER_CM", "Corer internal diameter (cm)")):
        ref = f"'1. Instructions'!$B${loi_rows[key]}"
        dn = DefinedName(name, attr_text=ref)
        try:
            wb.defined_names[name] = dn
        except TypeError:
            wb.defined_names.append(dn)

    if example:
        example(wb, ins, log, smp, loi_rows)
    return wb


# ----------------------------------------------------------------------------- example
def cowichan(wb, ins, log, smp, loi_rows):
    """Cowichan Estuary eelgrass cores from Janousek et al. (2025) — real published values."""
    cores = {r["SampID"]: r for r in csv.DictReader(open("data/reference/janousek2025_zostera_cores.csv"))
             if r["Estuary"] == "COW"}
    depth = [r for r in csv.DictReader(open("data/reference/janousek2025_zostera_depthseries.csv"))
             if r["SampID"] in cores and r["BD_type"] == "M"]
    dia = 7.6  # Douglas et al. (2022): acrylic tubes, 7.6 cm diameter
    area = math.pi * (dia / 2) ** 2
    ins.cell(loi_rows["Corer internal diameter (cm)"], 2).value = dia
    ins.cell(loi_rows["LOI intercept"], 2).value = -0.197
    ins.cell(loi_rows["LOI slope"], 2).value = 0.320
    ins.cell(loi_rows["LOI intercept"], 3).value = ("Example: local calibration from the 16 slices in these cores "
                                                     "with both elemental OC and LOI (r² = 0.73).")
    r = N_CORE_ROWS[0]
    for sid in sorted(cores):
        c = cores[sid]
        rows = [d for d in depth if d["SampID"] == sid]
        code = rows[0]["StudySampID"]
        vals = [code, f"COW-{code}", None, None, "Cowichan Estuary, BC — eelgrass",
                float(c["Lat"]), float(c["Long"]), None, None, None, None, None]  # diameter: Instructions
        for i, v in enumerate(vals, start=1):
            log.cell(r, i).value = v
        log.cell(r, 16).value = "SG"
        log.cell(r, 17).value = "assume none"
        log.cell(r, 18).value = ("Published core (Douglas et al. 2022, via Janousek et al. 2025). Sampling date "
                                 "and insertion depth not published; source reports no compaction observed.")
        r += 1
    r = N_SLICE_ROWS[0]
    for sid in sorted(cores):
        rows = sorted((d for d in depth if d["SampID"] == sid), key=lambda d: float(d["depth_top_cm"]))
        code = rows[0]["StudySampID"]
        for k, d in enumerate(rows, start=1):
            top, bot, bd = float(d["depth_top_cm"]), float(d["depth_bottom_cm"]), float(d["BD"])
            if d["C_type"] == "M" and d["PercC"] not in ("", "NA"):
                val, kind, note = float(d["PercC"]), "OC", "OC: elemental analyser, carbonate removed by acid fumigation"
            elif d["PercOM"] not in ("", "NA"):
                val, kind, note = float(d["PercOM"]), "LOI", "LOI 550 °C, 5 h (no elemental C for this slice)"
            else:
                val, kind, note = None, None, "No carbon value in the source dataset"
            smp.cell(r, 1).value = f"COW-{code}"
            smp.cell(r, 2).value = k
            smp.cell(r, 3).value = top
            smp.cell(r, 4).value = bot
            smp.cell(r, 5).value = note + ". Dry weight back-calculated from published bulk density."
            smp.cell(r, 13).value = round(bd * area * (bot - top), 2)
            smp.cell(r, 14).value = val
            smp.cell(r, 15).value = kind
            r += 1


# ----------------------------------------------------------------------------- synthetic
# A computer-generated stratified survey, used ONLY to show how Option B's stratified estimate
# and interval work. It is not field data and is placed at 0° N, 0° E so it describes no real
# place. Three strata of unequal area; the third is deliberately left unsampled.
SYN_DIR = "data/synthetic"
SYN_SEED = 2026
SYN_LOI = (-0.197, 0.320)
DEG = 6371008.8 * math.pi / 180  # metres per degree near 0° N — same Earth radius as R polygon_area_m2()
SYN_STRATA = [  # name, lon from, lon to (degrees); all span lat 0.0000–0.0024
    ("dense", 0.0000, 0.0012), ("sparse", 0.0012, 0.0019), ("channel_edge", 0.0019, 0.0021)]
SYN_LAT = (0.0000, 0.0024)
SYN_PLOTS = {"dense": 6, "sparse": 5}          # channel_edge: none (unsampled on purpose)
SYN_SHAPE = {"dense": (0.45, 0.95, 12.0, 1.30), "sparse": (0.30, 0.45, 10.0, 1.50)}  # floor, excess, scale cm, BD


def syn_area_m2(lon0, lon1):
    return (lon1 - lon0) * DEG * (SYN_LAT[1] - SYN_LAT[0]) * DEG


def synthetic(wb, ins, log, smp, loi_rows):
    import random
    rng = random.Random(SYN_SEED)
    for ws in (ins, log, smp, wb["4. Core Summary"]):
        ws["A1"] = "SYNTHETIC DATA — not field observations · " + str(ws["A1"].value)
    dia = 7.0
    area = math.pi * (dia / 2) ** 2
    ins.cell(loi_rows["Corer internal diameter (cm)"], 2).value = dia
    ins.cell(loi_rows["LOI intercept"], 2).value, ins.cell(loi_rows["LOI slope"], 2).value = SYN_LOI
    ins.cell(loi_rows["LOI intercept"], 3).value = "SYNTHETIC: equation chosen for the demonstration, not fitted."
    cores = []
    for name, lon0, lon1 in SYN_STRATA:
        for k in range(1, SYN_PLOTS.get(name, 0) + 1):
            pid = f"SYN-{name[0].upper()}{k}"
            lon = rng.uniform(lon0 + 0.0001, lon1 - 0.0001)
            lat = rng.uniform(SYN_LAT[0] + 0.0001, SYN_LAT[1] - 0.0001)
            mult = math.exp(rng.gauss(0, 0.25))
            n_cores = 2 if (name, k) == ("dense", 3) else 1      # one plot with two cores
            for j in range(n_cores):
                cid = pid + ("ab"[j] if n_cores > 1 else "")
                cores.append((name, pid, cid, lon + 0.00003 * j, lat, mult))
    r = N_CORE_ROWS[0]; sr = N_SLICE_ROWS[0]
    for i, (name, pid, cid, lon, lat, mult) in enumerate(cores):
        extracted = rng.choice([30, 32, 34, 36, 38, 40, 42])
        inserted = round(extracted * rng.uniform(1.05, 1.18))
        vals = [pid, cid, "2026-07-15", None, "SYNTHETIC site", round(lat, 6), round(lon, 6), None, None,
                None, inserted, extracted]
        for c, v in enumerate(vals, start=1):
            log.cell(r, c).value = v
        log.cell(r, 16).value = name
        log.cell(r, 18).value = f"SYNTHETIC — generated by data-raw/build_workbooks.py (seed {SYN_SEED})."
        r += 1
        floor, excess, scale, bd0 = SYN_SHAPE[name]
        use_loi = i % 4 == 1
        edges = [0, 2, 4, 6, 8, 10] + list(range(15, extracted, 5)) + [extracted]
        edges = sorted(set(e for e in edges if e <= extracted))
        cf = inserted / extracted
        for k, (top, bot) in enumerate(zip(edges[:-1], edges[1:]), start=1):
            mid = (top + bot) / 2 * cf
            oc = max(0.05, mult * (floor + excess * math.exp(-mid / scale)) * math.exp(rng.gauss(0, 0.08)))
            bd = (bd0 + 0.006 * mid) * math.exp(rng.gauss(0, 0.04)) * cf   # in-tube (compressed) density
            smp.cell(sr, 1).value = cid
            smp.cell(sr, 2).value = k
            smp.cell(sr, 3).value = top
            smp.cell(sr, 4).value = bot
            smp.cell(sr, 5).value = "SYNTHETIC"
            smp.cell(sr, 13).value = round(bd * area * (bot - top), 2)
            if use_loi:
                smp.cell(sr, 14).value = round((oc - SYN_LOI[0]) / SYN_LOI[1], 2)
                smp.cell(sr, 15).value = "LOI"
            else:
                smp.cell(sr, 14).value = round(oc, 3)
                smp.cell(sr, 15).value = "OC"
            sr += 1


def write_synthetic_area_files():
    import os
    os.makedirs(SYN_DIR, exist_ok=True)
    with open(f"{SYN_DIR}/SYNTHETIC_boundary.csv", "w", newline="") as f:
        w = csv.writer(f); w.writerow(["longitude", "latitude"])
        lon0, lon1 = SYN_STRATA[0][1], SYN_STRATA[-1][2]
        for x, y in [(lon0, SYN_LAT[0]), (lon1, SYN_LAT[0]), (lon1, SYN_LAT[1]), (lon0, SYN_LAT[1])]:
            w.writerow([f"{x:.4f}", f"{y:.4f}"])
    with open(f"{SYN_DIR}/SYNTHETIC_strata.csv", "w", newline="") as f:
        w = csv.writer(f); w.writerow(["stratum", "longitude", "latitude"])
        for name, lon0, lon1 in SYN_STRATA:
            for x, y in [(lon0, SYN_LAT[0]), (lon1, SYN_LAT[0]), (lon1, SYN_LAT[1]), (lon0, SYN_LAT[1])]:
                w.writerow([name, f"{x:.4f}", f"{y:.4f}"])
    return {name: round(syn_area_m2(lon0, lon1)) for name, lon0, lon1 in SYN_STRATA}


if __name__ == "__main__":
    build().save(f"{OUT_DIR}/Eelgrass_Carbon_DigitalData_BlankSheet.xlsx")
    build(cowichan).save(f"{OUT_DIR}/Eelgrass_Carbon_DigitalData_Example.xlsx")
    areas = write_synthetic_area_files()
    build(synthetic).save(f"{SYN_DIR}/Eelgrass_Carbon_DigitalData_SYNTHETIC.xlsx")
    print("written; synthetic stratum areas (m2):", areas)
